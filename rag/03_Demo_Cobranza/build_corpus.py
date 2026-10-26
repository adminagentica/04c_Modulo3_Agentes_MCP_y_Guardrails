"""Construye el corpus docente versionado. El linaje real lo registra la ingesta."""
import argparse, csv, hashlib, html, json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parent
DATE='2026-10-20'
def digest(data): return hashlib.sha256(data).hexdigest()
def canon(d): return json.dumps(d,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode('utf-8')
def table(doc, rows, widths):
    t=doc.add_table(rows=0,cols=len(widths));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,values in enumerate(rows):
        row=t.add_row()
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'153650' if i==0 else ('F0F4F7' if i%2==0 else 'FFFFFF'));tc.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tc.append(borders);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
            tc.append(m);cell.text=str(value)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
                for r in p.runs:r.font.size=Pt(9);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '172033')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def make_doc(d,out):
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
    for name,size in [('Normal',10),('Title',24),('Heading 1',15),('Heading 2',11)]:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(5)
        if name!='Normal':st.paragraph_format.keep_with_next=True
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    for st in [doc.styles['Title'],doc.styles['Heading 1'],doc.styles['Heading 2']]:
        for border in list(st._element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(6)
    doc.add_paragraph(d['title'],'Title');doc.add_paragraph(f"{d['code']} · {d['version']} · Corpus sintético de clase")
    doc.add_paragraph('Aprobación experta simulada para la clase. Las áreas y el revisor de este documento son ficticios; no constituye una política corporativa real.')
    table(doc,[('Campo','Valor'),('Propietario',d['owner']),('Aprobador de referencia',d['approver']),('Estado de uso','Vigente en el ejercicio docente'),('Publicación y aplicación',DATE),('Revisor simulado',d['reviewer']),('Próxima revisión',d['review_due']),('Reemplaza',f"{d['code']} {d['supersedes']} — conservado en histórico"),('Acceso',d['access']),('Ubicación del original',d['source_path'])],[1.85,4.9])
    doc.add_heading('Propósito',1);doc.add_paragraph(d['purpose'])
    doc.add_heading('Alcance',1);doc.add_paragraph(d['scope'])
    doc.add_page_break();doc.add_heading('Reglas y procedimientos',1)
    for title,text in d['sections']:
        doc.add_heading(title,2);doc.add_paragraph(text)
    # Delimita el contenido normativo que se extrae para el RAG.
    doc.add_heading('Controles del sistema RAG',1)
    doc.add_paragraph('Aplicar permisos y vigencia antes de recuperar. Conservar documento, versión, sección y huella del archivo en cada cita. Abstenerse si falta evidencia; las reglas de pausa y la aprobación humana se comprueban en la aplicación.')
    doc.add_page_break();doc.add_heading('Definiciones y referencias',1)
    for term,definition in d['glossary']:
        p=doc.add_paragraph();p.add_run(term+'. ').bold=True;p.add_run(definition)
    doc.add_paragraph('Las referencias siguientes identifican las versiones aplicables en este corpus. La sección se indica en cada regla. Una contradicción obliga a pausar y consultar al responsable; no se elige una regla por similitud de búsqueda.')
    table(doc,[('Documento relacionado','Versión')]+[(r['code'],r['version']) for r in d['references']],[5.45,1.3])
    doc.add_heading('Control y revisión documental',1)
    doc.add_paragraph(f"Revisión y aprobación simuladas: {d['reviewer']}, {DATE}. Alcance: estructura, terminología, coherencia entre las cuatro fuentes y correspondencia de los casos de Ground Truth. No acredita validación jurídica ni aprobación de una organización.")
    doc.add_paragraph(f"El propietario revisa esta versión el {d['review_due']} o antes si cambian reglas, fuentes o permisos. Una modificación genera una nueva versión, archiva la anterior y obliga a reindexar y revalidar los casos afectados. Se conservará la evidencia de revisión.")
    doc.add_heading('Historial de versiones',1)
    table(doc,[('Versión','Fecha','Cambio')]+[tuple(x) for x in d['history']],[.6,1,5.15])
    doc.add_heading('Trazabilidad de la fuente',1)
    doc.add_paragraph(f"Original: {d['source_path']}. Versión reemplazada: {d['previous_path']}. El mapa de secciones está en lineage_manifest.json; el registro de procesamiento real está en processing_lineage.json. Los parámetros del índice se mantienen fuera de esta política.")
    p=doc.add_paragraph('SHA-256 canónico: '+d['hash']);p.paragraph_format.space_after=Pt(3)
    for r in p.runs:r.font.size=Pt(8)
    doc.add_paragraph('El hash canónico identifica el contenido estructurado del documento según corpus_spec.json, serializado como JSON UTF-8 con claves ordenadas y sin espacios separadores. La huella del archivo Word completo se calcula después de guardarlo y se registra por separado en el manifiesto. No son valores intercambiables.')
    for p in sec.footer.paragraphs:
        p.text=f"{d['code']} {d['version']} · Aprobación simulada · Uso docente"
        for r in p.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
    filename=f"{d['code']}_{d['version']}.docx";doc.save(out/filename)
    return filename

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--output',type=Path,default=ROOT.parent/'05_Corpus_Documental');ap.add_argument('--web',type=Path);args=ap.parse_args()
    args.output.mkdir(parents=True,exist_ok=True);spec=json.loads((ROOT/'corpus_spec.json').read_text());items=[]
    for source in spec:
        d={**source,'hash':digest(canon(source))};filename=make_doc(d,args.output)
        items.append({**d,'filename':filename,'file_hash':digest((args.output/filename).read_bytes()),'canonical_schema':'corpus_spec.json / JSON sorted compact UTF-8'})
    (args.output/'lineage_manifest.json').write_text(json.dumps({'generated_at':DATE,'classification':'Datos sintéticos para formación','approval_mode':'simulated','documents':items},ensure_ascii=False,indent=2)+'\n')
    with (args.output/'lineage_manifest.csv').open('w',newline='') as f:
        w=csv.writer(f);w.writerow(['document_code','version','valid_from','review_due','approval_mode','canonical_sha256','file_sha256','chunk_count'])
        for d in items:w.writerow([d['code'],d['version'],d['valid_from'],d['review_due'],'simulated',d['hash'],d['file_hash'],len(d['chunks'])])
    if args.web:
        args.web.mkdir(parents=True,exist_ok=True)
        for d in items:
            body=''.join('<h2>'+html.escape(t)+'</h2><p>'+html.escape(s)+'</p>' for t,s in d['sections'])
            page=f"<!doctype html><html lang='es'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><link rel='stylesheet' href='docs.css'><title>{html.escape(d['title'])}</title><main><a href='/'>Volver a la demo</a><h1>{html.escape(d['title'])}</h1><p>{d['code']} {d['version']} · Aprobación experta simulada para la clase · {DATE}</p><p>Fuente sintética. Descargar el <a href='/api/rag/files/{d['code']}'>Word original</a> o ver <a href='/rag/sources/{d['code']}'>la evidencia realmente indexada</a>.</p>{body}<h2>Linaje de procesamiento</h2><p>Embeddings embeddinggemma:latest, 768 dimensiones. SQLite FTS5 y coseno exacto. El registro real se genera tras la ingesta en processing_lineage.json.</p></main></html>"
            (args.web/(d['code'].lower()+'.html')).write_text(page)
        (args.web/'index.html').write_text("<!doctype html><html lang='es'><meta charset='utf-8'><link rel='stylesheet' href='docs.css'><title>Corpus documental</title><main><a href='/'>Volver a la demo</a><h1>Corpus documental de clase</h1><p>Aprobación experta simulada. Fuentes sintéticas.</p>"+''.join(f"<section><a href='{d['code'].lower()}.html'>{html.escape(d['title'])} · {d['version']}</a><p>{len(d['chunks'])} secciones · <a href='/rag/sources/{d['code']}'>Ver fuente indexada</a></p></section>" for d in items)+"</main></html>")
    print(json.dumps({'documents':[d['filename'] for d in items],'sections':sum(len(d['chunks']) for d in items)},ensure_ascii=False))
if __name__=='__main__':main()
